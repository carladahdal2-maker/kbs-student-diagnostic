import io
from datetime import datetime

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

FONT = "Arial"
GREEN = RGBColor(0x1B, 0x6B, 0x3A)      

_STATUS_AR = {"confirmed": "مؤكَّد", "likely": "مرجَّح", "weak": "ضعيف"}
_PRIORITY_AR = {"critical": "حرجة", "high": "عالية", "medium": "متوسطة", "low": "منخفضة"}
_ERR_KIND_AR = {
    "conceptual": "خلل فهمي",
    "procedural": "خلل تطبيقي",
    "mixed": "مختلط (فهمي وتطبيقي)",
    "careless": "زلّات (إهمال)",
}


# ══════════════════════════════════════════════════════════════ #
#                       أدوات التنسيق (RTL + Arial)             #
# ══════════════════════════════════════════════════════════════ #
def _apply_font(run, size, bold, color):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.rtl = True
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:hAnsi"), FONT)
    rfonts.set(qn("w:cs"), FONT)


def _rtl_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.append(OxmlElement("w:bidi"))
    paragraph.alignment = align


def _add_para(container, text, size=11, bold=False, color=None,
              align=WD_ALIGN_PARAGRAPH.RIGHT):
    p = container.add_paragraph()
    _rtl_paragraph(p, align)
    _apply_font(p.add_run(text), size, bold, color)
    return p


def _fill_cell(cell, text, bold=False, color=None, size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    _rtl_paragraph(p)
    _apply_font(p.add_run(str(text)), size, bold, color)


def _new_table(doc, headers):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table._tbl.tblPr.append(OxmlElement("w:bidiVisual"))   # جدول RTL
    for cell, head in zip(table.rows[0].cells, headers):
        _fill_cell(cell, head, bold=True, color=GREEN)
    return table


# ══════════════════════════════════════════════════════════════ #
#                       اشتقاق نوع الخلل                        #
# ══════════════════════════════════════════════════════════════ #
def _error_kind(topic, diag_by_topic):
    """نوع الخلل من التشخيص، أو احتياطياً من عدّادات الموضوع."""
    kind = diag_by_topic.get(topic["topic_id"])
    if kind in _ERR_KIND_AR:
        return _ERR_KIND_AR[kind]
    c = topic.get("conceptual_errors", 0)
    a = topic.get("application_errors", 0)
    k = topic.get("careless_errors", 0)
    if c + a + k == 0:
        return "—"
    if k > c + a:
        return _ERR_KIND_AR["careless"]
    if c > a:
        return _ERR_KIND_AR["conceptual"]
    if a > c:
        return _ERR_KIND_AR["procedural"]
    return _ERR_KIND_AR["mixed"]


# ══════════════════════════════════════════════════════════════ #
#                       الدالة الرئيسية                         #
# ══════════════════════════════════════════════════════════════ #
def generate_report(result: dict, output_path: str = None) -> bytes:
    """تبني تقرير .docx من ناتج interpret() وتُرجع محتواه bytes.

    إن مُرّر output_path حُفظ الملف هناك أيضاً. تعمل مع أي نتيجة حتى الخالية.
    """
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    _rfonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    _rfonts.set(qn("w:ascii"), FONT)
    _rfonts.set(qn("w:hAnsi"), FONT)
    _rfonts.set(qn("w:cs"), FONT)

    student_id = result.get("student_id") or "—"
    quiz_id = result.get("quiz_id") or "—"

    # ── العنوان ──────────────────────────────────────────── #
    _add_para(doc, "تقرير أداء الطالب", size=20, bold=True, color=GREEN,
              align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_para(doc, f"الطالب: {student_id}  —  الاختبار: {quiz_id}",
              size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    # ── المستوى والنتيجة ─────────────────────────────────── #
    _add_para(doc, "الملخّص العام", size=15, bold=True, color=GREEN)
    _add_para(doc, f"المستوى العام: {result.get('level') or '—'}", size=12, bold=True)
    _add_para(doc, f"النتيجة المئوية: {result.get('score', 0)}%", size=12, bold=True)

    # ── نقاط القوة ───────────────────────────────────────── #
    _add_para(doc, "نقاط القوة", size=15, bold=True, color=GREEN)
    strengths = result.get("strengths", [])
    if strengths:
        table = _new_table(doc, ["الموضوع", "الدقة %"])
        for t in strengths:
            row = table.add_row().cells
            _fill_cell(row[0], t["topic_name"])
            _fill_cell(row[1], f"{t['accuracy']}%")
    else:
        _add_para(doc, "لا توجد نقاط قوة مسجّلة.", size=11)

    # ── نقاط الضعف ───────────────────────────────────────── #
    _add_para(doc, "نقاط الضعف", size=15, bold=True, color=GREEN)
    weaknesses = result.get("weaknesses", [])
    if weaknesses:
        diag_by_topic = {}
        for d in result.get("_diagnoses", []):
            if d["kind"] in _ERR_KIND_AR:
                diag_by_topic[d["topic"]] = d["kind"]
        table = _new_table(doc, ["الموضوع", "الدقة %", "نوع الخلل"])
        for t in weaknesses:
            row = table.add_row().cells
            _fill_cell(row[0], t["topic_name"])
            _fill_cell(row[1], f"{t['accuracy']}%")
            _fill_cell(row[2], _error_kind(t, diag_by_topic))
    else:
        _add_para(doc, "لا توجد نقاط ضعف مسجّلة.", size=11)

    # ── المفاهيم الخاطئة المكتشفة ────────────────────────── #
    misconceptions = result.get("misconceptions", [])
    if misconceptions:
        _add_para(doc, "المفاهيم الخاطئة المكتشفة", size=15, bold=True, color=GREEN)
        table = _new_table(doc, ["المفهوم الخاطئ", "الحالة", "الثقة", "الأسئلة الشاهدة"])
        for m in misconceptions:
            row = table.add_row().cells
            _fill_cell(row[0], m["label"])
            _fill_cell(row[1], _STATUS_AR.get(m["status"], m["status"]))
            _fill_cell(row[2], m["cf"])
            _fill_cell(row[3], "، ".join(m.get("occurrences", [])) or "—")

    # ── التوصيات ─────────────────────────────────────────── #
    _add_para(doc, "التوصيات الدراسية", size=15, bold=True, color=GREEN)
    recommendations = result.get("recommendations", [])
    if recommendations:
        for r in recommendations:
            prio = _PRIORITY_AR.get(r["priority"], r["priority"])
            _add_para(doc, f"[{prio}] {r['message']}", size=12, bold=True)
            _add_para(doc, f"↳ {r['explanation']}", size=11)
    else:
        _add_para(doc, "لا توجد توصيات.", size=11)

    # ── تاريخ الإنشاء ────────────────────────────────────── #
    doc.add_paragraph()
    _add_para(doc, f"تاريخ إنشاء التقرير: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
              size=9, color=RGBColor(0x77, 0x77, 0x77))

    # ── الإخراج ──────────────────────────────────────────── #
    buffer = io.BytesIO()
    doc.save(buffer)
    data = buffer.getvalue()
    if output_path:
        with open(output_path, "wb") as f:
            f.write(data)
    return data
